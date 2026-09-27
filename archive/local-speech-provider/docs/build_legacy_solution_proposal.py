from __future__ import annotations

from pathlib import Path
from typing import Iterable

from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
ASSET_DIR = ROOT / "docs" / "assets" / "方案文档"
OUTPUT = ROOT / "docs" / "Beyond-Words产品与技术落地方案.docx"

BLUE = "2563EB"
DEEP_BLUE = "163D77"
PALE_BLUE = "EFF6FF"
PALE_GRAY = "F8FAFC"
MID_GRAY = "64748B"
LIGHT_GRAY = "D9D9D9"
BLACK = "000000"
WHITE = "FFFFFF"


def set_run_font(run, name: str = "Microsoft YaHei", size: float | None = None, bold: bool | None = None,
                 color: str | None = None):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), name)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color is not None:
        run.font.color.rgb = RGBColor.from_string(color)
    return run


def set_cell_shading(cell, fill: str):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_border(cell, color: str = LIGHT_GRAY, size: str = "4"):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = f"w:{edge}"
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def set_cell_margins(cell, top: int = 100, start: int = 120, bottom: int = 100, end: int = 120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_row_cant_split(row):
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    cant_split.set(qn("w:val"), "true")
    tr_pr.append(cant_split)


def set_paragraph_spacing(paragraph, before: float = 0, after: float = 6, line: float = 1.22):
    fmt = paragraph.paragraph_format
    fmt.space_before = Pt(before)
    fmt.space_after = Pt(after)
    fmt.line_spacing = line
    fmt.widow_control = True


def add_body(doc: Document, text: str, bold_lead: str | None = None, after: float = 6):
    p = doc.add_paragraph()
    set_paragraph_spacing(p, after=after)
    if bold_lead and text.startswith(bold_lead):
        set_run_font(p.add_run(bold_lead), size=10.5, bold=True, color=BLACK)
        set_run_font(p.add_run(text[len(bold_lead):]), size=10.5, color="1F2937")
    else:
        set_run_font(p.add_run(text), size=10.5, color="1F2937")
    return p


def add_bullets(doc: Document, items: Iterable[str], level: int = 0):
    for item in items:
        p = doc.add_paragraph(style="List Bullet" if level == 0 else "List Bullet 2")
        set_paragraph_spacing(p, after=3, line=1.16)
        set_run_font(p.add_run(item), size=10.2, color="1F2937")


def add_numbered(doc: Document, items: Iterable[str]):
    for index, item in enumerate(items, start=1):
        p = doc.add_paragraph()
        set_paragraph_spacing(p, after=4, line=1.18)
        p.paragraph_format.left_indent = Inches(0.24)
        p.paragraph_format.first_line_indent = Inches(-0.24)
        set_run_font(p.add_run(f"{index}.  "), size=10.2, bold=True, color="334155")
        set_run_font(p.add_run(item), size=10.2, color="1F2937")


def add_hyperlink(paragraph, text: str, url: str):
    part = paragraph.part
    rel_id = part.relate_to(url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True)
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), rel_id)
    run = OxmlElement("w:r")
    r_pr = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), BLUE)
    r_pr.append(color)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    r_pr.append(underline)
    r_fonts = OxmlElement("w:rFonts")
    r_fonts.set(qn("w:ascii"), "Microsoft YaHei")
    r_fonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    r_pr.append(r_fonts)
    run.append(r_pr)
    text_node = OxmlElement("w:t")
    text_node.text = text
    run.append(text_node)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def add_code_block(doc: Document, text: str):
    p = doc.add_paragraph()
    set_paragraph_spacing(p, before=5, after=8, line=1.0)
    p.paragraph_format.left_indent = Inches(0.22)
    p.paragraph_format.right_indent = Inches(0.22)
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(8)
    p_pr = p._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), "F1F5F9")
    p_pr.append(shd)
    set_run_font(p.add_run(text), name="Consolas", size=9.5, color="0F172A")
    return p


def set_alt_text(picture_run, title: str, description: str):
    drawings = picture_run._element.xpath(".//wp:docPr")
    if drawings:
        drawings[0].set("title", title)
        drawings[0].set("descr", description)


def add_figure(doc: Document, path: Path, caption: str, width: float = 6.4, alt: str | None = None):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_paragraph_spacing(p, before=4, after=3, line=1.0)
    run = p.add_run()
    run.add_picture(str(path), width=Inches(width))
    set_alt_text(run, caption, alt or caption)
    cap = doc.add_paragraph()
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_paragraph_spacing(cap, after=8, line=1.0)
    set_run_font(cap.add_run(caption), size=8.5, color=MID_GRAY)


def style_data_table(table, widths: list[float] | None = None):
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    for row_idx, row in enumerate(table.rows):
        for col_idx, cell in enumerate(row.cells):
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_border(cell)
            set_cell_margins(cell)
            if widths:
                cell.width = Inches(widths[col_idx])
            if row_idx == 0:
                set_cell_shading(cell, DEEP_BLUE)
            elif row_idx % 2 == 0:
                set_cell_shading(cell, PALE_BLUE)
            else:
                set_cell_shading(cell, WHITE)
            for p in cell.paragraphs:
                set_paragraph_spacing(p, after=0, line=1.1)
                for run in p.runs:
                    set_run_font(run, size=9, bold=row_idx == 0, color=WHITE if row_idx == 0 else "1F2937")
    set_repeat_table_header(table.rows[0])


def add_table(doc: Document, headers: list[str], rows: list[list[str]], widths: list[float]):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    for idx, header in enumerate(headers):
        table.rows[0].cells[idx].text = header
    for row_data in rows:
        cells = table.add_row().cells
        for idx, value in enumerate(row_data):
            cells[idx].text = value
    style_data_table(table, widths)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return table


def add_picture_cell(cell, image_path: Path, caption: str, notes: list[str], width: float = 2.35):
    cell.text = ""
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
    set_cell_margins(cell, top=90, start=90, bottom=100, end=90)
    set_cell_border(cell)
    picture_p = cell.paragraphs[0]
    picture_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_paragraph_spacing(picture_p, after=4, line=1.0)
    picture_run = picture_p.add_run()
    picture_run.add_picture(str(image_path), width=Inches(width))
    set_alt_text(picture_run, caption, caption)
    caption_p = cell.add_paragraph()
    caption_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_paragraph_spacing(caption_p, after=5, line=1.0)
    set_run_font(caption_p.add_run(caption), size=8.5, bold=True, color="334155")
    for note in notes:
        p = cell.add_paragraph(style="List Bullet")
        set_paragraph_spacing(p, after=2, line=1.08)
        set_run_font(p.add_run(note), size=8.2, color="475569")


def add_picture_pair(doc: Document, left: tuple, right: tuple):
    table = doc.add_table(rows=1, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    add_picture_cell(table.cell(0, 0), *left)
    add_picture_cell(table.cell(0, 1), *right)
    set_row_cant_split(table.rows[0])
    set_repeat_table_header(table.rows[0])
    doc.add_paragraph().paragraph_format.space_after = Pt(1)


def add_heading(doc: Document, text: str, level: int = 1):
    p = doc.add_heading(text, level=level)
    p.paragraph_format.keep_with_next = True
    if level == 1 and text != "方案结论":
        p.paragraph_format.page_break_before = True
    return p


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run()
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = " PAGE "
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.append(fld_char1)
    run._r.append(instr_text)
    run._r.append(fld_char2)
    set_run_font(run, size=8, color=MID_GRAY)


def font(size: int, bold: bool = False):
    candidates = [
        Path(r"C:\Windows\Fonts\msyhbd.ttc" if bold else r"C:\Windows\Fonts\msyh.ttc"),
        Path(r"C:\Windows\Fonts\simhei.ttf"),
        Path(r"C:\Windows\Fonts\arial.ttf"),
    ]
    for candidate in candidates:
        if candidate.exists():
            return ImageFont.truetype(str(candidate), size=size)
    return ImageFont.load_default()


def multiline_center(draw: ImageDraw.ImageDraw, box: tuple[int, int, int, int], lines: list[str], fnt, fill: str,
                     gap: int = 10):
    x1, y1, x2, y2 = box
    sizes = [draw.textbbox((0, 0), line, font=fnt) for line in lines]
    heights = [bbox[3] - bbox[1] for bbox in sizes]
    total = sum(heights) + gap * (len(lines) - 1)
    y = y1 + (y2 - y1 - total) / 2
    for line, bbox, height in zip(lines, sizes, heights):
        width = bbox[2] - bbox[0]
        draw.text((x1 + (x2 - x1 - width) / 2, y), line, font=fnt, fill=fill)
        y += height + gap


def draw_arrow(draw: ImageDraw.ImageDraw, start: tuple[int, int], end: tuple[int, int], color: str = "#94A3B8"):
    draw.line([start, end], fill=color, width=8)
    ex, ey = end
    draw.polygon([(ex, ey), (ex - 22, ey - 14), (ex - 22, ey + 14)], fill=color)


def create_brand_mark(path: Path):
    image = Image.new("RGBA", (360, 360), (255, 255, 255, 0))
    draw = ImageDraw.Draw(image)
    draw.rounded_rectangle((20, 20, 340, 340), radius=88, fill="#2563EB")
    centers = [125, 165, 205, 245]
    heights = [95, 175, 125, 70]
    for x, height in zip(centers, heights):
        draw.rounded_rectangle((x - 9, 180 - height // 2, x + 9, 180 + height // 2), radius=9, fill="#FFFFFF")
    image.save(path)


def create_user_flow(path: Path):
    image = Image.new("RGB", (1800, 640), "#F8FAFC")
    draw = ImageDraw.Draw(image)
    title_font = font(54, bold=True)
    node_font = font(34, bold=True)
    small_font = font(25)
    draw.text((70, 45), "核心用户闭环", font=title_font, fill="#0F172A")
    labels = [
        ("取得同意", "选择场景"),
        ("真实录音", "标记时刻"),
        ("自动处理", "转写与分人"),
        ("查看复盘", "核对证据"),
        ("完成练习", "再次对话"),
    ]
    node_w, node_h, gap = 270, 250, 62
    x0, y0 = 70, 210
    fills = ["#EFF6FF", "#E0EAFF", "#DBEAFE", "#EDE9FE", "#DCFCE7"]
    strokes = ["#93C5FD", "#60A5FA", "#3B82F6", "#A78BFA", "#86EFAC"]
    for idx, ((primary, secondary), fill_color, stroke) in enumerate(zip(labels, fills, strokes)):
        x1 = x0 + idx * (node_w + gap)
        box = (x1, y0, x1 + node_w, y0 + node_h)
        draw.rounded_rectangle(box, radius=38, fill=fill_color, outline=stroke, width=5)
        draw.ellipse((x1 + 102, y0 + 30, x1 + 168, y0 + 96), fill="#2563EB")
        number = str(idx + 1)
        nb = draw.textbbox((0, 0), number, font=small_font)
        draw.text((x1 + 135 - (nb[2] - nb[0]) / 2, y0 + 48), number, font=small_font, fill="#FFFFFF")
        multiline_center(draw, (x1 + 15, y0 + 110, x1 + node_w - 15, y0 + node_h - 22), [primary, secondary], node_font, "#0F172A", 12)
        if idx < len(labels) - 1:
            draw_arrow(draw, (x1 + node_w + 8, y0 + node_h // 2), (x1 + node_w + gap - 12, y0 + node_h // 2))
    image.save(path, quality=95)


def create_architecture(path: Path):
    image = Image.new("RGB", (1800, 940), "#FFFFFF")
    draw = ImageDraw.Draw(image)
    title_font = font(54, bold=True)
    group_font = font(30, bold=True)
    item_font = font(25)
    draw.text((70, 48), "本地优先的技术架构", font=title_font, fill="#0F172A")

    groups = [
        (90, 190, 430, 720, "用户端", ["React + TypeScript", "Zustand", "MediaRecorder", "IndexedDB"]),
        (535, 190, 900, 720, "应用后端", ["FastAPI", "REST API", "SQLite", "本地音频目录"]),
        (1005, 190, 1370, 720, "语音与语义", ["Whisper 英文转写", "pyannote 双人分离", "本地规则", "可选兼容大模型"]),
        (1475, 190, 1735, 720, "输出", ["复盘", "证据片段", "互动洞察", "定向练习"]),
    ]
    colors = [("#EFF6FF", "#60A5FA"), ("#F8FAFC", "#94A3B8"), ("#F5F3FF", "#A78BFA"), ("#F0FDF4", "#4ADE80")]
    for idx, (x1, y1, x2, y2, label, items) in enumerate(groups):
        fill_color, outline = colors[idx]
        draw.rounded_rectangle((x1, y1, x2, y2), radius=40, fill=fill_color, outline=outline, width=5)
        label_box = draw.textbbox((0, 0), label, font=group_font)
        draw.text(((x1 + x2 - (label_box[2] - label_box[0])) / 2, y1 + 44), label, font=group_font, fill="#0F172A")
        y = y1 + 135
        for item in items:
            draw.rounded_rectangle((x1 + 35, y, x2 - 35, y + 72), radius=20, fill="#FFFFFF", outline="#CBD5E1", width=3)
            bbox = draw.textbbox((0, 0), item, font=item_font)
            draw.text(((x1 + x2 - (bbox[2] - bbox[0])) / 2, y + 20), item, font=item_font, fill="#334155")
            y += 92
        if idx < len(groups) - 1:
            draw_arrow(draw, (x2 + 16, 455), (groups[idx + 1][0] - 16, 455), "#64748B")

    draw.rounded_rectangle((225, 785, 1575, 885), radius=28, fill="#0F172A")
    footer = "默认录音不外发  只有显式配置外部语义模型时才发送匿名转写文字"
    bbox = draw.textbbox((0, 0), footer, font=group_font)
    draw.text(((1800 - (bbox[2] - bbox[0])) / 2, 817), footer, font=group_font, fill="#FFFFFF")
    image.save(path, quality=95)


def configure_document(doc: Document):
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.68)
    section.bottom_margin = Inches(0.65)
    section.left_margin = Inches(0.72)
    section.right_margin = Inches(0.72)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Microsoft YaHei"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor.from_string("1F2937")

    title = styles["Title"]
    title.font.name = "Microsoft YaHei"
    title._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    title.font.size = Pt(28)
    title.font.bold = True
    title.font.color.rgb = RGBColor.from_string(BLACK)
    title.paragraph_format.space_after = Pt(12)

    for style_name, size, before, after in (("Heading 1", 18, 14, 8), ("Heading 2", 13.5, 10, 5), ("Heading 3", 11.5, 8, 4)):
        style = styles[style_name]
        style.font.name = "Microsoft YaHei"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(BLACK)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    for section in doc.sections:
        footer = section.footer
        footer_p = footer.paragraphs[0]
        set_run_font(footer_p.add_run("Beyond Words 产品与技术落地方案  |  "), size=8, color=MID_GRAY)
        add_page_number(footer_p)


def build_document():
    ASSET_DIR.mkdir(parents=True, exist_ok=True)
    brand_path = ASSET_DIR / "brand-mark.png"
    flow_path = ASSET_DIR / "用户流程图.png"
    architecture_path = ASSET_DIR / "技术架构图.png"
    create_brand_mark(brand_path)
    create_user_flow(flow_path)
    create_architecture(architecture_path)

    doc = Document()
    configure_document(doc)
    doc.core_properties.title = "Beyond Words 产品与技术落地方案"
    doc.core_properties.subject = "真实英语对话复盘应用的产品界面技术架构模型和落地计划"
    doc.core_properties.author = "Beyond Words 项目组"
    doc.core_properties.keywords = "英语对话, 语音转写, 说话人分离, 互动分析, MVP"

    # Cover
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_paragraph_spacing(p, before=18, after=18, line=1.0)
    brand_run = p.add_run()
    brand_run.add_picture(str(brand_path), width=Inches(0.92))
    set_alt_text(brand_run, "Beyond Words 标志", "蓝色圆角方形中的白色声波图形")
    title = doc.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_run_font(title.add_run("Beyond Words 产品与技术落地方案"), size=28, bold=True, color=BLACK)
    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_paragraph_spacing(subtitle, after=10)
    set_run_font(subtitle.add_run("真实英语对话复盘  自动转写  双人区分  证据化洞察  定向练习"), size=12, color=MID_GRAY)
    meta = doc.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_paragraph_spacing(meta, after=30)
    set_run_font(meta.add_run("版本 1.0    2026 年 9 月 25 日    用于 Pro 账号续作与团队评审"), size=9.5, color="94A3B8")

    add_heading(doc, "方案结论", 1)
    add_body(doc, "当前代码已经具备可运行的移动端界面、浏览器录音、本地持久化、FastAPI 后端、真实 Whisper 英文转写、复盘、互动时刻和练习闭环。后续工作的核心不是重写界面，而是补齐双人说话人分离、提升转写质量、建立浏览器端到端测试，并把分析任务迁移到可恢复的后台工作进程。")
    add_body(doc, "产品定位：Beyond Words 不给用户打综合分，也不推断人格或情绪；它把原始录音、时间戳话轮和可回听证据组织成可核对的交流复盘，再把具体互动时刻转成下一轮可以练习的表达。", bold_lead="产品定位：")
    add_table(doc, ["当前状态", "结论"], [
        ["界面与主流程", "可演示并可继续开发"],
        ["真实英文转写", "本机 Whisper tiny.en 已验证"],
        ["双人自动区分", "接口已预留，当前依赖未安装"],
        ["语义分析", "本地规则可用，外部大模型可选"],
        ["实体 soundcore", "尚未接 SDK，界面明确标为模拟"],
    ], [1.65, 4.65])

    # Plan mode
    add_heading(doc, "如何开启 Codex 计划模式", 1)
    add_body(doc, "在 Codex 输入框中直接输入下面的斜杠命令并发送，即可切换计划模式。官方说明指出，计划模式会先给出可审阅的实施路径，并在需要时提出澄清问题，再开始实际改动。")
    add_code_block(doc, "/plan")
    add_heading(doc, "适合使用计划模式的情况", 2)
    add_bullets(doc, [
        "任务会同时改动前端、后端、模型、数据结构或部署方式。",
        "存在隐私、成本、模型选择或兼容性风险，需要先定边界。",
        "你希望先审阅工作清单、验收标准和停止点，再允许实现。",
        "需要让另一个账号或开发者按同一份方案继续工作。",
    ])
    add_heading(doc, "推荐的首条计划提示", 2)
    add_code_block(doc, "请阅读项目总纲、README、docs/当前验收记录与后续工作.md 和本方案文档。只输出可执行计划，不修改代码。计划必须列出 P0/P1 工作包、受影响文件、依赖、数据风险、验收方式和明确停止点。")
    add_heading(doc, "确认计划后的执行方式", 2)
    add_numbered(doc, [
        "审阅计划，删掉不需要的范围，并补充真实设备、模型账号或部署环境约束。",
        "切回默认执行模式，或直接回复“按已确认计划执行 P0 第一项”。",
        "每完成一个工作包就运行测试、更新进度文档，并保留可回滚结果。",
        "涉及麦克风授权、外部 API 密钥、删除数据或付费服务时，由你本人完成确认。",
    ])
    p = doc.add_paragraph()
    set_paragraph_spacing(p, before=8, after=6)
    set_run_font(p.add_run("官方来源："), size=9, bold=True, color="334155")
    add_hyperlink(p, "Run long horizon tasks with Codex", "https://developers.openai.com/blog/run-long-horizon-tasks-with-codex")

    # Product scope
    add_heading(doc, "产品目标与范围", 1)
    add_heading(doc, "目标用户", 2)
    add_body(doc, "主要用户是需要在真实英语交流后复盘参与方式的学习者，包括英语角参与者、留学生、小组讨论成员、课堂学习者和面试准备者。产品不要求他们在对话中盯着屏幕，而是在对话结束后提供可解释的回顾。")
    add_heading(doc, "核心问题", 2)
    add_table(doc, ["用户问题", "产品回应", "证据"], [
        ["不知道自己是否真正参与了对话", "展示发言时长、占比和话轮", "带时间戳的话轮"],
        ["只知道自己说错了词，不知道互动哪里卡住", "识别话题延续、追问、澄清等时刻", "前后话轮与原声片段"],
        ["看到建议但不会在下次使用", "把时刻转成文字或语音练习", "提示、回答和反馈记录"],
        ["担心录音和模型外发", "默认本机处理并明确披露数据流", "隐私页与导出删除能力"],
    ], [1.8, 2.45, 2.05])
    add_heading(doc, "必须保留的产品边界", 2)
    add_bullets(doc, [
        "不根据一次对话推断人格、动机、情绪、心理状态或综合英语能力。",
        "所有自动洞察都绑定原始话轮和时间范围，并提示用户核对。",
        "未配置模型时清楚显示能力降级，不用演示结果冒充真实推断。",
        "录音前必须取得所有参与者的知情同意。",
        "实体硬件能力未接入时保持“模拟设备”标识。",
    ])

    add_figure(doc, flow_path, "图 1 核心用户闭环", 6.5, "从知情同意开始，经录音、自动处理、复盘和练习回到下一次真实对话的五步闭环")

    # Visual system
    add_heading(doc, "界面设计方案", 1)
    add_body(doc, "现有界面采用移动端优先布局，最大内容宽度 430 像素。视觉层次以深蓝、亮蓝、白色和浅灰为主，录音页切换到深色沉浸式环境；复盘和练习页面保持浅色，帮助用户阅读长文本和对比证据。")
    add_table(doc, ["设计元素", "当前做法", "使用原则"], [
        ["主色", "亮蓝 #2563EB", "主行动、选中状态、用户话轮"],
        ["深色", "深蓝黑 #0B1220", "只用于录音等沉浸场景"],
        ["状态色", "绿、黄、红", "分别表达可查看、提示与破坏性操作"],
        ["卡片", "16 至 28 像素圆角", "用于分组，不堆叠无意义阴影"],
        ["文字", "中文优先，英文保留原话", "标题简短，说明明确，避免给用户贴标签"],
        ["导航", "底部四入口", "首页、复盘、练习、我的保持稳定"],
    ], [1.25, 2.0, 3.05])
    add_heading(doc, "交互原则", 2)
    add_bullets(doc, [
        "录音是显式动作：先同意，再请求麦克风权限，再开始采集。",
        "分析结果先给摘要，再给指标、时间轴、逐句转写和重点时刻，阅读顺序从概览到证据。",
        "不确定性要可见：模拟分析、真实转写、说话人待确认和模型来源分别标注。",
        "练习入口紧跟具体互动时刻，避免用户离开语境后再寻找练习。",
        "删除、外部传输和权限请求要让用户明确确认。",
    ])

    # Screens 1-2
    add_heading(doc, "首页与录音界面", 1)
    add_body(doc, "首页负责发起一次对话并提供轻量进度反馈；录音页负责把知情同意、场景、音量、时间和标记动作集中在一个沉浸界面中。")
    add_picture_pair(doc,
        (ASSET_DIR / "01-首页.png", "图 2 首页", ["主行动足够突出", "进度卡片不使用总分", "最近会话直接进入复盘"]),
        (ASSET_DIR / "02-录音.png", "图 3 录音页", ["深色降低环境干扰", "录音豆承担开始和停止", "同意勾选位于权限动作之前"]),
    )
    add_heading(doc, "实现说明", 2)
    add_body(doc, "首页由 React 路由加载 Zustand 中的本地会话，并与后端会话合并。录音页使用 MediaRecorder 采集 WebM 音频，先写入浏览器 IndexedDB，再同步到 FastAPI；同步失败时保留本地副本和重试入口。单击开始或停止，录音中双击或点击按钮添加重点标记。")

    # Screens 3-4
    add_heading(doc, "历史与复盘概览", 1)
    add_body(doc, "历史页承担检索和删除入口；复盘概览先回答“这段对话发生了什么”，再展示发言指标和话轮时间轴。真实数据与演示数据使用不同披露条。")
    add_picture_pair(doc,
        (ASSET_DIR / "03-历史复盘.png", "图 4 历史复盘", ["标题和场景可搜索", "真实会话显示删除入口", "演示会话保留为体验样例"]),
        (ASSET_DIR / "04-复盘概览.png", "图 5 复盘概览", ["摘要先于指标", "时间轴区分你与伙伴", "模拟分析有醒目标识"]),
    )
    add_heading(doc, "自动区分完成后的界面变化", 2)
    add_body(doc, "当 pyannote 输出两个匿名说话人后，复盘页应先让用户回听样例并确认“哪位说话人是我”。只有确认完成后才计算发言占比、用户话轮数和针对用户的互动事件，避免把伙伴的话误算给用户。")

    # Screens 5-6
    add_heading(doc, "转写核对与重点时刻", 1)
    add_body(doc, "自动分析必须能被用户核对和修正。逐句话轮保留时间范围、说话人和原始文字；重点时刻把建议拆成发生了什么、语境解释和可尝试的另一种表达。")
    add_heading(doc, "事件类型", 2)
    add_table(doc, ["类型", "识别信号", "用户得到的建议"], [
        ["话题延续", "简短回答后对方重新提问", "补充理由并邀请对方参与"],
        ["主动追问", "对方提供信息后话题结束", "表达兴趣并提出开放问题"],
        ["请求澄清", "重复、停顿或答非所问", "明确指出需要澄清的部分"],
        ["小组参与", "较长空档或多轮未参与", "选择已有观点并自然接入"],
    ], [1.25, 2.45, 2.6])
    add_heading(doc, "核对闭环", 2)
    add_numbered(doc, [
        "从重点时刻回放对应音频，同时展示前后话轮和系统依据。",
        "用户先修正转写，再修正匿名说话人归属，所有指标随之重算。",
        "用户确认建议是否符合语境；不合适的结果可忽略，不写回能力标签。",
        "确认后的时刻才能进入定向练习，练习记录保留来源会话和模型版本。",
    ])
    add_heading(doc, "界面示意", 2)
    add_picture_pair(doc,
        (ASSET_DIR / "05-转写与洞察.png", "图 6 转写核对", ["按话轮显示时间戳", "真实会话支持修改文字", "双人模型可用后支持修正说话人"], 2.55),
        (ASSET_DIR / "06-重点时刻.png", "图 7 重点时刻", ["保留原声片段", "建议绑定相关话轮", "明确表示解释而非评分"], 2.55),
    )

    # Screens 7-8
    add_heading(doc, "定向练习与个人目标", 1)
    add_body(doc, "练习页让用户在相同语境中重说一轮，而不是做脱离对话的通用题；个人页记录交流目标和连续行为，不显示综合能力分数。")
    add_picture_pair(doc,
        (ASSET_DIR / "07-定向练习.png", "图 8 定向练习", ["复用真实对方话轮", "支持文字与语音回答", "反馈聚焦策略而非人格"]),
        (ASSET_DIR / "08-个人目标.png", "图 9 个人目标", ["选择一个当前交流目标", "展示会话和练习数量", "模拟硬件状态明确披露"]),
    )
    add_heading(doc, "练习反馈规则", 2)
    add_body(doc, "MVP 的本地反馈只检查是否表达观点、是否补充理由、是否提出能让对方继续分享的问题，并展示结构化建议。配置外部语义模型后可以改善自然度和上下文理解，但仍不得输出人格、情绪或综合能力判断。")

    # Data screen
    add_heading(doc, "隐私与数据界面", 1)
    add_body(doc, "隐私页把数据实际保存位置和外部传输条件写在操作按钮之前。导出用于可携带性，删除全部数据是破坏性操作，必须二次确认。")
    add_figure(doc, ASSET_DIR / "09-隐私与数据.png", "图 10 隐私与数据", 2.7, "界面说明录音保存在浏览器和本机后端，外部模型只接收匿名转写文字，并提供导出和删除按钮")
    add_table(doc, ["数据", "默认位置", "外部模型是否接收", "删除方式"], [
        ["原始录音", "浏览器 IndexedDB 和本机音频目录", "否", "删除会话或全部数据"],
        ["转写与时间戳", "SQLite 分析 JSON", "仅在主动配置后", "删除会话或全部数据"],
        ["说话人标签", "匿名 ID，不记录真实身份", "可随转写一起发送", "删除会话或全部数据"],
        ["练习回答", "SQLite 与浏览器状态", "视语义模型配置", "删除全部数据"],
    ], [1.15, 2.25, 1.55, 1.35])

    # Architecture
    add_heading(doc, "技术架构与数据流", 1)
    add_figure(doc, architecture_path, "图 11 本地优先技术架构", 6.55, "React 前端经 REST 调用 FastAPI，数据保存到 SQLite 与本地音频目录，语音模型和语义分析生成复盘与练习")
    add_heading(doc, "前端", 2)
    add_body(doc, "React 19、TypeScript 和 Vite 负责界面与构建；React Router 管理八类页面；Zustand 持久化会话、练习与个人目标；IndexedDB 保存浏览器端音频；MediaRecorder 与 Web Audio API 完成录音和实时音量。")
    add_heading(doc, "后端", 2)
    add_body(doc, "FastAPI 提供会话、音频、标记、分析、说话人确认、转写修正、练习、设置、导出和删除接口。SQLite 保存结构化记录，本地目录保存音频，分析任务具备状态记录和服务重启后的恢复逻辑。")
    add_heading(doc, "数据流", 2)
    add_numbered(doc, [
        "用户确认同意后开始录音，音频先写入 IndexedDB。",
        "前端创建远程会话、上传音频和标记，失败时保留本地状态。",
        "后端执行 ASR，再根据可用性执行双人分离和语义事件分析。",
        "用户确认自己的说话人身份，系统重新计算个人指标和事件。",
        "复盘和练习结果写入 SQLite，可导出或删除。",
    ])

    # Model stack
    add_heading(doc, "模型方案", 1)
    add_table(doc, ["能力", "推荐实现", "当前状态", "验收指标"], [
        ["英文语音转写", "Whisper base.en 或 small.en", "tiny.en 可运行", "WER、实时系数、显存"],
        ["双人说话人分离", "pyannote speaker-diarization-community-1", "依赖未安装", "DER、双人数稳定性"],
        ["说话人身份", "模型匿名分组 + 用户确认", "界面和接口已实现", "确认后指标一致"],
        ["互动事件", "本地证据规则 + 可选兼容 LLM", "本地规则可用", "证据命中、误报率、可解释性"],
        ["练习反馈", "结构规则 + 可选兼容 LLM", "基础规则可用", "建议一致性与延迟"],
    ], [1.25, 2.15, 1.35, 1.55])
    add_heading(doc, "为什么大模型不是必需入口", 2)
    add_body(doc, "录音、转写、时间戳和双人区分分别由浏览器能力、Whisper 和 pyannote 完成。通用大模型主要用于摘要、互动解释和练习反馈；没有配置时，本地规则仍能给出透明、可复现的基础结果。因此大模型应作为可插拔增强层，而不是让核心录音链路依赖外部服务。")
    add_heading(doc, "当前评测证据", 2)
    add_table(doc, ["样本", "模型", "结果", "判断"], [
        ["HCRC Map Task", "Whisper tiny.en", "WER 25.09%", "可用于流程验证"],
        ["AMI ES2002a", "Whisper tiny.en", "WER 41.51%", "多人会议质量不足"],
        ["AMI 说话人分离", "pyannote", "DER 尚无结果", "必须补跑"],
        ["现有 6 秒录音", "Whisper tiny.en", "约 7.1 秒完成", "Triton 未启用，性能待优化"],
    ], [1.55, 1.55, 1.45, 1.75])
    add_heading(doc, "模型选择规则", 2)
    add_bullets(doc, [
        "使用同一批 MapTask、AMI 和真实产品录音比较 tiny.en、base.en、small.en。",
        "同时记录 WER、DER、处理时长、峰值显存和失败样本，不只比较准确率。",
        "如果外部语义模型不可用或返回格式错误，必须自动退回本地规则。",
        "外部模型输入只包含匿名转写话轮和必要时间信息，不上传录音。",
    ])

    # Backend feature map
    add_heading(doc, "后端功能完成清单", 1)
    add_table(doc, ["功能域", "已有能力", "仍需补齐"], [
        ["会话与音频", "创建、上传、读取、重命名、删除", "分片上传、配额、格式校验"],
        ["分析任务", "持久状态、失败原因、重启恢复", "独立 worker、取消、超时和幂等"],
        ["语音模型", "Whisper 真实转写、时间戳、降级", "模型基准、双人分离、性能优化"],
        ["复盘", "摘要、话轮、指标、事件、转写修正", "真实双人闭环与质量监控"],
        ["练习", "文字和语音入口、转写、反馈、历史", "更可靠的语义反馈与版本记录"],
        ["隐私", "同意门槛、导出、删除、本地优先", "身份认证、日志脱敏、保留策略"],
        ["硬件", "模拟状态和浏览器麦克风", "soundcore SDK、连接和按键事件"],
    ], [1.25, 2.35, 2.75])
    add_heading(doc, "关键接口分类", 2)
    add_table(doc, ["接口组", "作用"], [
        ["/api/sessions", "会话、标题、场景和状态管理"],
        ["/api/sessions/{id}/audio", "音频上传和回放"],
        ["/api/sessions/{id}/analyze", "启动异步分析任务"],
        ["/api/sessions/{id}/speakers", "确认哪位说话人是用户"],
        ["/api/sessions/{id}/turns/{turnId}", "修正文字或说话人"],
        ["/api/practices 与 attempts", "创建练习并保存反馈"],
        ["/api/export 与 /api/data", "数据导出和全部删除"],
    ], [2.75, 3.55])

    # Execution plan
    add_heading(doc, "落地执行方案", 1)
    add_body(doc, "执行顺序以能否形成真实、可重复验收的产品闭环为依据。每个工作包都要有代码、测试、运行记录和文档更新，不能只完成安装或页面展示。")
    add_heading(doc, "P0 双人自动区分与模型基准", 2)
    add_numbered(doc, [
        "创建独立 Python 环境，固定 CUDA、PyTorch、Whisper 和 pyannote 版本。",
        "完成 Community-1 模型许可和 Hugging Face 令牌配置。",
        "用一段本地双人录音验证匿名说话人时间段与合并话轮。",
        "在 AMI 上计算 DER，在 MapTask 和 AMI 上对比三个 ASR 模型。",
        "选择默认模型并把配置、指标和硬件条件写入 baselines.json。",
        "完成“系统分为两人—用户确认自己—指标和事件更新”的界面闭环。",
    ])
    add_heading(doc, "P0 浏览器端到端测试", 2)
    add_numbered(doc, [
        "建立 Playwright 测试配置和独立测试数据库。",
        "使用虚拟音频输入覆盖同意、录音、标记、停止、同步和分析。",
        "覆盖后端断线、本地保存、恢复和重试上传。",
        "覆盖复盘、修改转写、说话人确认、练习、导出和删除确认。",
        "在持续集成中同时运行后端测试、前端测试、构建和 E2E。",
    ])
    add_heading(doc, "P0 版本与数据基线", 2)
    add_numbered(doc, [
        "初始化 Git 或迁移到明确的远程仓库，保存当前通过验收的基线提交。",
        "确认用户录音、运行数据库、模型文件、密钥、虚拟环境和原始大语料不进入 Git。",
        "把源码包和评测数据包分别保存，保证接手者能重建而不会拿到用户隐私数据。",
    ])
    add_heading(doc, "P1 生产化", 2)
    add_bullets(doc, [
        "将分析任务迁移到独立 worker 和持久队列，增加取消、超时、重试与进度心跳。",
        "增加认证、速率限制、上传大小和类型校验、数据库迁移与备份恢复。",
        "完成日志脱敏、保留期限、网络请求清单和录音不外发断言。",
        "有 SDK 和测试设备后实现 soundcore 适配层，同时保留浏览器麦克风降级路径。",
    ])

    # Acceptance and risks
    add_heading(doc, "验收方案", 1)
    add_table(doc, ["验收对象", "必须证明的结果", "证据"], [
        ["前端", "主要路由、交互、错误状态和移动端布局正常", "单元测试、E2E、截图"],
        ["后端", "接口、持久化、任务恢复和删除一致", "临时数据库测试与接口日志"],
        ["ASR", "真实录音可转写且质量可量化", "WER、延迟、失败样本"],
        ["说话人分离", "双人稳定分组且可人工纠正", "DER、实录回听、确认结果"],
        ["语义分析", "提示绑定证据，失败时安全降级", "事件样例、格式测试、故障注入"],
        ["隐私", "录音不外发，导出与删除完整", "网络审计、数据目录检查"],
        ["硬件", "真实连接和按键事件可恢复", "真机录像与断连测试"],
    ], [1.25, 3.0, 2.1])
    add_heading(doc, "当前风险", 2)
    add_table(doc, ["优先级", "风险", "处理方式"], [
        ["高", "说话人模型尚不可用", "先完成依赖和 DER，再开放双方指标"],
        ["高", "tiny.en 在 AMI 上 WER 41.51%", "比较 base.en 和 small.en"],
        ["高", "缺少完整 E2E", "用虚拟音频覆盖核心旅程"],
        ["中", "Whisper 未启用 Triton 快速内核", "固定兼容 CUDA 组合并测实时系数"],
        ["中", "分析仍在 Web 进程中", "拆分 worker 与任务队列"],
        ["中", "浏览器和后端双状态", "增加幂等同步和冲突测试"],
        ["低", "桌面宽屏留白", "发布桌面版前再决定双栏布局"],
    ], [0.8, 2.35, 3.2])
    add_heading(doc, "继续工作的起点", 2)
    add_body(doc, "下一个账号进入项目后，应先运行现有测试并阅读交接文档，然后从“P0 双人自动区分与模型基准”开始。不要先接外部大模型或美化桌面布局，因为这两项都不能解决当前最关键的真实性和质量问题。")
    add_body(doc, "当前测试基线：后端 8/8 通过，前端 5/5 通过，生产构建通过；真实本地 Whisper 转写已抽查成功。", bold_lead="当前测试基线：")
    add_body(doc, "停止原则：需要用户提供密钥、接受模型许可、授予麦克风权限、删除数据或连接实体设备时停止自动执行，记录当前状态并请求明确确认。", bold_lead="停止原则：")

    add_heading(doc, "项目内参考资料", 2)
    add_bullets(doc, [
        "README.zh-CN.md：启动方式、功能与环境配置。",
        "docs/项目进度与Pro账号续作交接.md：完整现状和接手信息。",
        "docs/当前验收记录与后续工作.md：最近验收证据与用户验证项。",
        "docs/技术架构与功能实现.md：代码地图、模型和接口说明。",
        "evaluation/README.zh-CN.md：真实语料导入与 WER/DER 评测。",
    ])
    add_heading(doc, "接手账号开工前检查", 2)
    add_bullets(doc, [
        "确认只保留源码、配置示例、文档、测试和必要的小型样例；不要提交 node_modules、虚拟环境、模型缓存、运行数据库或用户录音。",
        "先阅读中文 README、Pro 账号续作交接文档和最近验收记录，再运行现有测试与生产构建。",
        "核对前端 5173 端口和后端 8000 端口没有被旧进程占用，且前端确实连接到当前后端。",
        "需要 Hugging Face 令牌、模型许可、外部大模型密钥或真实设备时，记录所需信息并等待用户提供。",
        "每完成一个 P0 工作包就更新基准数据、验收证据和后续工作，不把尚未验证的能力写成已完成。",
    ])

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build_document()

