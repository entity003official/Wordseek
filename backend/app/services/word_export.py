from io import BytesIO
import re
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

def review_document(session, language="zh"):
    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Cm(21), Cm(29.7)
    section.top_margin = section.bottom_margin = Cm(2)
    section.left_margin = section.right_margin = Cm(2.2)
    for name in ["Normal", "Title", "Heading 1", "Heading 2"]:
        style = doc.styles[name]
        style.font.name = "Calibri"
        style.element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Yu Gothic" if language == "ja" else "Microsoft YaHei")
        style.font.color.rgb = RGBColor.from_string("202326")
    normal = doc.styles["Normal"]
    normal.font.size = Pt(11)
    normal.paragraph_format.line_spacing = 1.3
    normal.paragraph_format.space_after = Pt(7)
    for name, size in [("Title",22),("Heading 1",16),("Heading 2",12)]:
        doc.styles[name].font.size = Pt(size)
        doc.styles[name].paragraph_format.keep_with_next = True
    labels = {
        "zh": ["语言学习复盘", "对话概况", "重点词汇", "同义替换", "自然表达", "原句", "解释", "替代表达或搭配", "使用区别", "例句", "完整转写", "说话人", "本人", "尚未生成语言学习复盘", "录音时长", "来源话轮"],
        "en": ["Language learning review", "Overview", "Vocabulary", "Alternatives", "Natural expressions", "Original", "Meaning", "Alternative or phrase", "Usage", "Example", "Transcript", "Speaker", "You", "Language review has not been generated", "Duration", "Source turns"],
        "ja": ["言語学習の振り返り", "会話の概要", "重要な語彙", "類義表現", "自然な表現", "元の表現", "説明", "言い換えや組み合わせ", "使い分け", "例文", "文字起こし", "話者", "自分", "学習ポイントはまだ生成されていません", "録音時間", "引用元"]
    }.get(language)
    if labels is None: labels = ["Review"] * 16
    def clean(value):
        return re.sub(r"[\x00-\x08\x0b\x0c\x0e-\x1f]", "", str(value or ""))
    def paragraph(value, style=None):
        return doc.add_paragraph(clean(value), style)
    paragraph(session.title, "Title")
    seconds = max(0, session.duration_ms or 0) // 1000
    paragraph(f"{labels[0]} · {session.created_at.strftime('%Y-%m-%d')} · {labels[14]} {seconds//60:02d}:{seconds%60:02d}")
    analysis = session.analysis_json or {}
    review = analysis.get("ai_review") or {}
    if review:
        doc.add_heading(labels[1], level=1)
        paragraph(review.get("summary", {}).get("summary", ""))
        for kind, title in [("vocabulary",labels[2]),("synonym",labels[3]),("natural_expression",labels[4])]:
            points = [p for p in review.get("learning_points", []) if p.get("kind") == kind]
            if not points: continue
            doc.add_heading(title, level=1)
            for point in points:
                doc.add_heading(clean(point.get("original")), level=2)
                for field, label in [("explanation",labels[6]),("alternative",labels[7]),("usage_note",labels[8]),("example",labels[9])]:
                    paragraph(label + ": " + clean(point.get(field)))
                paragraph(labels[15] + ": " + ", ".join(point.get("evidence_turn_ids", [])))
        for note in review.get("summary", {}).get("limitations", []): paragraph(note)
    else:
        paragraph(labels[13])
    doc.add_heading(labels[10], level=1)
    ids = [speaker["id"] for speaker in analysis.get("speakers", [])]
    for turn in analysis.get("turns", []):
        speaker = turn.get("speaker_id")
        label = labels[12] if speaker == session.user_speaker_id else labels[11] + " " + str(ids.index(speaker)+1 if speaker in ids else "?")
        start = int(turn.get("start_ms",0)) // 1000
        end = int(turn.get("end_ms",0)) // 1000
        heading = paragraph(f"{start//60:02d}:{start%60:02d}–{end//60:02d}:{end%60:02d}  {label}")
        heading.paragraph_format.keep_with_next = True
        paragraph(turn.get("text"))
    footer = section.footer.paragraphs[0]
    footer.alignment = 2
    footer.add_run("Beyond Words  ·  ")
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    footer._p.append(field)
    output = BytesIO()
    doc.save(output)
    return output.getvalue()
